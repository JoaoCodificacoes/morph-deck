"""Build templates/neutral/neutral.pptx: a plain 16:9 template, free to share.

Starts from python-pptx's default (Office) template, widens it to 16:9, sets Aptos as the theme fonts, places
the title and slide number of the "Title Only" layout, and adds one thin accent rule above the title.
Run: python templates/neutral/make_template.py
"""
import re
from copy import deepcopy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from morphdeck.draw import box

OUT = Path(__file__).with_name("neutral.pptx")
W, H = 13.333, 7.5
COVER, CONTENT = 0, 5  # "Title Slide", "Title Only"


def widen(shapes, factor):
    for sh in shapes:
        if sh.has_text_frame or sh.is_placeholder:
            try:
                sh.left, sh.width = int(sh.left * factor), int(sh.width * factor)
            except (TypeError, AttributeError):  # inherits its position; nothing to scale
                pass


def align_left(ph):
    """Left-align a layout placeholder through its list style, which is what slides inherit."""
    body = ph._element.txBody
    lst = body.find(qn("a:lstStyle"))
    if lst is None:
        lst = etree.Element(qn("a:lstStyle"))
        body.find(qn("a:bodyPr")).addnext(lst)
    lvl = lst.find(qn("a:lvl1pPr"))
    if lvl is None:
        lvl = etree.SubElement(lst, qn("a:lvl1pPr"))
    lvl.set("algn", "l")


def main():
    prs = Presentation()
    factor = Inches(W) / prs.slide_width
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    widen(prs.slide_master.placeholders, factor)
    for lay in prs.slide_layouts:
        widen(lay.placeholders, factor)

    theme = next(p for p in prs.part.package.iter_parts() if "theme" in str(p.partname))
    xml = theme.blob.decode("utf8")
    xml = re.sub(r'(<a:majorFont><a:latin typeface=")[^"]*', r"\1Aptos Display", xml)
    xml = re.sub(r'(<a:minorFont><a:latin typeface=")[^"]*', r"\1Aptos", xml)
    theme._blob = xml.encode("utf8")

    content = prs.slide_layouts[CONTENT]
    for ph in content.placeholders:
        t = ph.placeholder_format.idx
        if t == 0:  # title
            ph.left, ph.top, ph.width, ph.height = Inches(0.7), Inches(0.62), Inches(11.9), Inches(0.85)
        elif t == 12:  # slide number
            ph.left, ph.top, ph.width, ph.height = Inches(11.9), Inches(6.85), Inches(0.7), Inches(0.35)
    cover = prs.slide_layouts[COVER]
    for lay in (content, cover):
        for ph in lay.placeholders:
            if ph.placeholder_format.idx in (0, 1):
                align_left(ph)
    for ph in cover.placeholders:
        if ph.placeholder_format.idx == 0:
            ph.left, ph.top, ph.width, ph.height = Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.6)
        elif ph.placeholder_format.idx == 1:
            ph.left, ph.top, ph.width, ph.height = Inches(0.9), Inches(4.1), Inches(11.5), Inches(1.2)

    # accent rule above the title, drawn on a scratch slide and moved into the layouts
    scratch = prs.slides.add_slide(prs.slide_layouts[6])
    for lay, (x, y) in ((content, (0.74, 0.45)), (cover, (0.94, 2.05))):
        rule = box(scratch, x, y, 0.55, 0.06, fill="45535E", radius=0)
        lay.shapes._spTree.append(deepcopy(rule._element))
    ids = prs.slides._sldIdLst
    prs.part.drop_rel(ids[0].rId)
    ids.remove(ids[0])
    prs.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
