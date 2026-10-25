"""Each lint rule fires on a deck that breaks it, and stays quiet on one that doesn't."""

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from morphdeck import animate, box, link, morph, path, stage, text
from morphdeck.lint import lint


def deck(n=2):
    prs = Presentation()
    slides = [prs.slides.add_slide(prs.slide_layouts[6]) for _ in range(n)]
    for s in slides[1:]:
        morph(s)
    return prs, slides


def rules(prs, tmp_path, level=None):
    p = tmp_path / "d.pptx"
    prs.save(p)
    return {f.rule for f in lint(p) if level is None or f.level == level}


def test_clean_deck_has_no_errors(tmp_path):
    prs, (a, b) = deck()
    link(box(a, 1, 1, 2, 1), "card")
    link(box(b, 5, 2, 3, 2), "card")
    extra = box(b, 1, 5, 1, 1)
    animate(b, [[stage([extra])]])
    assert rules(prs, tmp_path, "error") == set()


def test_morph_target_with_entrance(tmp_path):
    prs, (a, b) = deck()
    link(box(a, 1, 1, 2, 1), "card")
    target = link(box(b, 5, 2, 3, 2), "card")
    animate(b, [[stage([target])]])
    assert "morph-target-animated" in rules(prs, tmp_path)


def test_entrance_must_be_first(tmp_path):
    prs, (a, _) = deck()
    dot = box(a, 1, 1, 0.2, 0.2)
    animate(a, [[stage([dot], "motion", 600, path=path([(0, -1), (0, 0)])), stage([dot], delay=-600)]])
    assert "entrance-not-first" in rules(prs, tmp_path)


def test_entrance_first_is_fine(tmp_path):
    prs, (a, _) = deck()
    dot = box(a, 1, 1, 0.2, 0.2)
    animate(a, [[stage([dot], dur=120), stage([dot], "motion", 600, delay=-120, path=path([(0, -1), (0, 0)]))]])
    assert "entrance-not-first" not in rules(prs, tmp_path)


def test_motion_must_end_at_rest_when_morphed(tmp_path):
    prs, (a, b) = deck()
    cursor = link(box(a, 1, 1, 0.3, 0.3), "cursor")
    link(box(b, 1, 1, 0.3, 0.3), "cursor")
    animate(a, [[stage([cursor], "motion", 600, path=path([(0, 0), (2, 1)]))]])
    assert "motion-not-at-rest" in rules(prs, tmp_path)


def test_motion_then_exit_is_fine(tmp_path):
    prs, (a, b) = deck()
    ghost = link(box(a, 1, 1, 0.3, 0.3), "ghost")
    link(box(b, 1, 1, 0.3, 0.3), "ghost")
    animate(a, [[stage([ghost], "motion", 600, path=path([(0, 0), (2, 1)])), stage([ghost], "fade-out", 200)]])
    assert "motion-not-at-rest" not in rules(prs, tmp_path)


def test_duplicate_link(tmp_path):
    prs, (a,) = deck(1)
    link(box(a, 1, 1, 1, 1), "card")
    link(box(a, 3, 1, 1, 1), "card")
    assert "duplicate-link" in rules(prs, tmp_path)


def test_double_effect_list(tmp_path):
    prs, (a,) = deck(1)
    sh = box(a, 1, 1, 1, 1)
    sp = sh._element.spPr
    etree.SubElement(sp, qn("a:effectLst"))
    etree.SubElement(sp, qn("a:effectLst"))
    assert "double-effect-list" in rules(prs, tmp_path)


def test_inherited_font(tmp_path):
    prs, (a,) = deck(1)
    tb = a.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    tb.text_frame.text = "no font set"
    assert "inherited-font" in rules(prs, tmp_path)
    prs2, (b,) = deck(1)
    text(b, 1, 1, 3, 1, "explicit font")
    assert "inherited-font" not in rules(prs2, tmp_path)


def test_no_link_warning(tmp_path):
    prs, (a, b) = deck()
    box(a, 1, 1, 1, 1)
    box(b, 1, 1, 1, 1)
    assert "no-link" in rules(prs, tmp_path, "warning")


def test_missing_target(tmp_path):
    prs, (a,) = deck(1)
    sh = box(a, 1, 1, 1, 1)
    animate(a, [[stage([sh])]])
    sh._element.getparent().remove(sh._element)
    assert "missing-target" in rules(prs, tmp_path)
