"""The building blocks produce what PowerPoint expects."""
from pptx import Presentation

from morphdeck import animate, box, line, morph, path, stage, twin
from morphdeck.motion import WIPE


def test_path_is_relative_to_slide_size():
    assert path([(13.333, 0), (0, 0)]) == "M 1.0000 0.0000 L 0.0000 0.0000 E"
    assert path([(0, 0), (1, 0), (2, 0), (3, 0)], curve=True).startswith("M 0.0000 0.0000 C ")


def test_twin_gets_new_id_and_name():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    a = box(s, 1, 1, 2, 1, name="!!a")
    b = twin(s, a, "b")
    assert b.name == "!!b" and b.shape_id != a.shape_id
    assert (b.left, b.top, b.width, b.height) == (a.left, a.top, a.width, a.height)


def test_dash_comes_before_arrow_head():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    ln = line(s, 0, 0, 1, 1, arrow=True, dash=True).line._get_or_add_ln()
    tags = [c.tag.split("}")[1] for c in ln]
    assert tags.index("prstDash") < tags.index("tailEnd")


def test_animate_writes_one_click_per_entry(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    morph(s)
    a, b = box(s, 1, 1, 1, 1), box(s, 3, 1, 1, 1)
    animate(s, [[stage([a])], [stage([b], "wipe-right")]])
    xml = s._element.xml
    assert xml.count('delay="indefinite"') == 2
    assert f'presetSubtype="{WIPE["right"][0]}"' in xml
    prs.save(tmp_path / "x.pptx")  # and the result still saves
