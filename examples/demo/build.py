"""Demo deck: a book-lending platform for a neighbourhood library (fictional).

Every pattern of morph-deck in nine slides. Run from the repo root:

    python examples/demo/build.py          # writes examples/demo/demo.pptx and lints it
"""
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from morphdeck import (Deck, animate, box, click_build, fade, icon, label, line, logo, path, pill, polyline, rgb,
                       shadow, stage, text, to_front, twin)

HERE = Path(__file__).parent
deck = Deck(HERE / "demo.toml")
P = deck.palette
A = deck.area
REQ, DATA, APP = P.identity("requirements"), P.identity("data"), P.identity("app")


def kicker(s, words, tone=None):
    """Small label above the title: which part of the talk this slide belongs to."""
    text(s, A.x + 0.05, 0.2, 6, 0.3, words.upper(), size=11, bold=True, color=tone.dark if tone else P.muted)


# 1 ------------------------------------------------------------------ cover
s = deck.slide("cover", "Book Lending Platform for a Neighbourhood Library", morph=False, title_size=40)
deck.fill(s, 1, "Demo deck for morph-deck · fictional content", size=16, color=P.muted)

# 2 ------------------------------------------------------------------ the situation: two sources
s = deck.slide("content", "Lending today: two sources", morph=False)
fade(s)  # nothing on the cover carries over: a plain fade, not a fake link
kicker(s, "Context")
sheet = box(s, A.x, 1.85, 3.4, 3.5, fill=P.bad.tint, radius=0.05, name="!!src-sheet")
sheet_t = [icon(s, "table_view", A.x + 0.3, 2.1, 0.55, color=P.bad.dark),
           text(s, A.x + 0.3, 2.8, 3.0, 2.4, [[("Spreadsheet", {"bold": True, "size": 22, "color": P.bad.dark})],
                                                "", "one file per year", "copied by hand",
                                                "no history of reminders"], size=14, spacing=4)]
plus = text(s, A.x + 3.5, 3.3, 0.5, 0.6, "+", size=34, bold=True, color=P.muted, align="c")


def paper(x, rot, title, rows, caption, name):
    """A paper form: the page travels by Morph, so it stays outside the group holding its contents."""
    page = box(s, x, 1.9, 3.3, 3.4, fill=P.white, line="C9CFD6", radius=0.01, name=name)
    page.rotation = rot
    shadow(page)
    g = s.shapes.add_group_shape()
    text(g, x + 0.2, 2.05, 2.9, 0.35, title, size=13, bold=True)
    for k, row in enumerate(rows):
        y = 2.6 + k * 0.42
        line(g, x + 0.2, y + 0.36, x + 3.1, y + 0.36, color=P.line, width=0.75)
        text(g, x + 0.25, y, 2.8, 0.35, row, size=10.5, color="2B5CAB", font="Segoe Print")
    g.rotation = rot
    cap = text(s, x, 5.45, 3.3, 0.4, caption, size=11.5, color=P.muted, align="c")
    return [page, g, cap]


cards = paper(4.8, -2.0, "Ficha de empréstimo", ["Ana R. · 12/03", "O Principezinho", "devolvido 26/03"],
              "Loan card: one per book", "!!src-card")
wait = paper(8.55, 1.5, "Lista de espera", ["1. Rui M. — Saramago", "2. Ana R. — Mia Couto", "3. —"],
             "Waiting list: on the desk", "!!src-wait")
animate(s, [[stage([sheet] + sheet_t)], [stage([plus] + cards), stage(wait)]])

# 3 ------------------------------------------------------------------ contributions: each source becomes a card
s = deck.slide("content", "Contributions")
kicker(s, "Context")
contrib = [("01", "Requirements", "What a loan, a reservation and a reminder must record.", REQ, "plagiarism",
            "!!src-card"),
           ("02", "Data model", "One schema for books, members and loans, with the rules in the database.", DATA,
            "schema", "!!src-sheet"),
           ("03", "Lending app", "Phone app for the desk, reservations included, working offline.", APP,
            "phone_android", "!!src-wait")]
builds = []
for i, (num, head, body, tone, ic, src) in enumerate(contrib):
    x = A.x + i * 4.0
    target = box(s, x, 1.85, 3.8, 3.5, radius=0.05, name=src)  # static: the Morph target of a source on slide 2
    builds.append([box(s, x, 1.85, 3.8, 0.09, fill=tone.base, radius=0),
                   text(s, x + 0.3, 2.1, 1.5, 0.9, num, size=44, bold=True, color=tone.dark, font=deck.theme.font_head),
                   box(s, x + 2.9, 2.2, 0.62, 0.62, fill=tone.tint, radius=0.5),
                   icon(s, ic, x + 3.01, 2.31, 0.4, color=tone.dark),
                   text(s, x + 0.3, 3.15, 3.2, 0.5, head, size=19, bold=True),
                   text(s, x + 0.3, 3.75, 3.2, 1.4, body, size=13.5, color=P.muted)])
twin(s, target, "app")  # the last card also becomes the app box on the next slide: one source, two targets
click_build(s, builds)

# 4 ------------------------------------------------------------------ architecture, one component per click
s = deck.slide("content", "Architecture")
kicker(s, "Lending app", APP)
app = box(s, A.x, 2.7, 2.6, 1.5, fill=P.white, line=P.edge, radius=0.08, name="!!app")  # static: from card 03
label(app, [[("Phone app", {"bold": True, "size": 15})], [("at the desk", {"size": 11, "color": P.muted})]])
api = box(s, 4.35, 2.7, 3.6, 1.5, fill=P.white, line=P.edge, radius=0.08, name="!!api")
api_t = text(s, 4.35, 3.15, 3.6, 0.6, "Backend", size=15, bold=True, align="c", anchor="m", name="!!api-t")
db = box(s, 9.4, 2.7, 3.2, 1.5, fill=DATA.tint, radius=0.08, name="!!db")
label(db, [[("Database", {"bold": True, "size": 15, "color": DATA.dark})],
           [("books, members, loans", {"size": 11, "color": DATA.dark})]])
a1 = line(s, A.x + 2.6, 3.45, 4.35, 3.45, color=P.edge, width=1.75, arrow=True, name="!!a1")
a2 = line(s, 7.95, 3.45, 9.4, 3.45, color=P.edge, width=1.75, arrow=True, name="!!a2")
badge = logo(s, "postgresql", 12.35, 2.48)
animate(s, [[stage([a1], "wipe-right", 350), stage([api, api_t], delay=-100)],
            [stage([a2], "wipe-right", 350), stage([db], delay=-100)],
            [stage(badge)]])

# 5 ------------------------------------------------------------------ zoom into the backend (Morph)
s = deck.slide("content", "Every loan is checked")
kicker(s, "Lending app", APP)
box(s, A.x, 2.7, 2.6, 1.5, fill=P.white, line=P.edge, radius=0.08, name="!!app")  # same place: stays put
label(s.shapes[-1], [[("Phone app", {"bold": True, "size": 15})], [("at the desk", {"size": 11, "color": P.muted})]])
box(s, 4.0, 1.85, 4.3, 4.3, fill=P.white, line=P.edge, radius=0.05, name="!!api")  # widens
text(s, 4.25, 1.95, 3.0, 0.5, "Backend", size=15, bold=True, anchor="m", name="!!api-t")  # label glides along
box(s, 9.4, 2.7, 3.2, 1.5, fill=DATA.tint, radius=0.08, name="!!db")
label(s.shapes[-1], [[("Database", {"bold": True, "size": 15, "color": DATA.dark})],
                     [("books, members, loans", {"size": 11, "color": DATA.dark})]])
line(s, A.x + 2.6, 3.45, 4.0, 3.45, color=P.edge, width=1.75, arrow=True, name="!!a1")
line(s, 8.3, 3.45, 9.4, 3.45, color=P.edge, width=1.75, arrow=True, name="!!a2")
checks = []
for k, (ic, t) in enumerate([("login", "Signed in"), ("admin_panel_settings", "Allowed to lend"),
                             ("menu_book", "Book is on the shelf")]):
    y = 2.6 + k * 0.8
    checks.append([icon(s, ic, 4.3, y + 0.05, 0.36, color=P.slate),
                   text(s, 4.8, y, 3.3, 0.45, t, size=14, bold=True, anchor="m")])
refused = pill(s, 4.3, 5.1, 3.7, 0.5, "Lent twice? Refused by the database", fill=P.fail.tint, color=P.fail.dark,
               size=11.5)
click_build(s, checks + [[refused]])

# 6 ------------------------------------------------------------------ offline timeline, a dot carries each step
s = deck.slide("content", "Returning a book without network")
kicker(s, "Lending app", APP)
card = box(s, A.x, 1.85, 2.8, 4.3, fill=APP.tint, radius=0.06, name="!!app")  # the app box becomes this card
icon(s, "phone_android", A.x + 0.9, 2.5, 1.0, color=APP.dark)
text(s, A.x, 3.7, 2.8, 0.8, [[("Phone app", {"bold": True, "size": 16, "color": APP.dark})],
                              [("at the desk", {"size": 12, "color": APP.dark})]], align="c")
steps = [("cloud_off", "No network", "The return is recorded on the phone.", P.bad.base),
         ("save", "Saved on the phone", "It waits in an outbox, in order.", P.slate),
         ("sync", "Connection returns", "The outbox is sent; rejected records are shown.", P.good.base),
         ("block", "Sent twice? Refused", "The database refuses a second copy.", P.fail.base)]
TX = 4.3
groups = []
for k, (ic, t, body, col) in enumerate(steps):
    y = 1.85 + k * 1.1
    groups.append([box(s, TX - 0.25, y, 0.5, 0.5, fill=col, radius=0.5), icon(s, ic, TX - 0.15, y + 0.1, 0.3, color=P.white),
                   text(s, TX + 0.5, y - 0.05, 7.5, 0.9, [[(t, {"bold": True, "size": 15})],
                                                          [(body, {"size": 12.5, "color": P.muted})]])])
clicks = [[stage(groups[0])]]
for k in range(1, len(steps)):
    top, bottom = 1.85 + (k - 1) * 1.1 + 0.5, 1.85 + k * 1.1
    seg = line(s, TX, top, TX, bottom, color=P.line, width=2)
    dot = box(s, TX - 0.07, bottom + 0.1, 0.14, 0.14, fill=P.slate, radius=0.5)  # rests under the next step
    clicks.append([stage([seg], "wipe-down", 600),
                   stage([dot], dur=120, delay=-600),  # the entrance comes first, or the dot shows early
                   stage([dot], "motion", 600, delay=-120, path=path([(0, top - bottom - 0.17), (0, 0)])),
                   stage([dot], "fade-out", 200, delay=-120),
                   stage(groups[k], dur=350, delay=-150)])
    to_front(s, groups[k][:2])
animate(s, clicks)

# 7 ------------------------------------------------------------------ measured against a limit
s = deck.slide("content", "Response times", morph=False)
fade(s)  # a new part of the talk: nothing carries over
kicker(s, "Evaluation")
TX0, TW = 3.6, 5.0
LIM = TX0 + TW * 0.8
rows = [("Open a member", 0.42, 0.6, "0.42 s"), ("Record a loan", 0.31, 0.6, "0.31 s"),
        ("Search the catalogue", 0.48, 0.6, "0.48 s")]
bars = []
for k, (name, v, lim, shown) in enumerate(rows):
    y = 2.3 + k * 0.85
    text(s, A.x, y, 2.8, 0.42, name, size=13, bold=True, anchor="m")
    box(s, TX0, y + 0.05, TW, 0.32, fill="EEF1F4", radius=0.5)
    bars.append([box(s, TX0, y + 0.05, TW * 0.8 * v / lim, 0.32, fill=P.good.base, radius=0.5),
                 text(s, TX0 + TW * 0.8 * v / lim + 0.08, y, 1.0, 0.42, shown, size=12, bold=True, anchor="m")])
    text(s, LIM + 0.08, y, 0.8, 0.42, "0.6 s", size=10, color=P.muted, anchor="m")
END = 2.3 + 3 * 0.85 - 0.1
line(s, LIM, 2.15, LIM, END, color=P.ink, width=1.5, dash=True, name="!!limit")
text(s, LIM - 1.0, END + 0.04, 2.0, 0.3, "target", size=10.5, bold=True, align="c", name="!!limit-t")
animate(s, [[stage([b[0]], "wipe-right", 500), stage([b[1]], dur=250)] for b in bars])

# 8 ------------------------------------------------------------------ the limit line turns into the chart's limit
s = deck.slide("content", "Under load")
kicker(s, "Evaluation")
PX, PY, PW, PH = A.x + 0.7, 2.0, 7.0, 3.8
xy = lambda x, y: (PX + PW * x / 60, PY + PH * (1 - y / 1.2))
grid = [line(s, PX, xy(0, v)[1], PX + PW, xy(0, v)[1], color="E6EAEE", width=0.75) for v in (0, 0.4, 0.8, 1.2)]
grid += [text(s, PX - 0.6, xy(0, v)[1] - 0.15, 0.55, 0.3, f"{v:g} s", size=10, color=P.muted, align="r")
         for v in (0, 0.4, 0.8, 1.2)]
grid += [text(s, xy(v, 0)[0] - 0.3, PY + PH + 0.05, 0.6, 0.3, str(v), size=10, color=P.muted, align="c")
         for v in (0, 20, 40, 60)]
grid.append(text(s, PX, PY + PH + 0.32, PW, 0.3, "loans per minute", size=10.5, color=P.muted, align="c"))
line(s, PX, xy(0, 0.6)[1], PX + PW, xy(0, 0.6)[1], color=P.ink, width=1.5, dash=True, name="!!limit")  # static
text(s, PX + 0.05, xy(0, 0.6)[1] - 0.32, 2.0, 0.3, "target", size=10.5, bold=True, name="!!limit-t")
median = polyline(s, [xy(*p) for p in [(5, 0.12), (15, 0.13), (25, 0.15), (35, 0.18), (45, 0.24), (55, 0.4)]],
                  color="9FB3C4", width=2)
p95 = polyline(s, [xy(*p) for p in [(5, 0.2), (15, 0.22), (25, 0.26), (35, 0.33), (45, 0.45), (55, 0.82)]],
               color=P.navy, width=3)
legend = text(s, PX + 0.05, PY + 0.05, 2.5, 0.5, [[("— slowest 5%", {"bold": True, "color": P.navy, "size": 10.5})],
                                                  [("— median", {"bold": True, "color": "8A9DB0", "size": 10.5})]])
stat = box(s, 9.2, 2.0, 3.4, 1.4, fill=P.card)
stat_t = text(s, 9.4, 2.1, 3.1, 1.2, [[("45 / min", {"bold": True, "size": 26, "color": P.good.dark})],
                                       [("within target, no errors", {"size": 12})]])
animate(s, [[stage(grid)],
            [stage([median], "wipe-right", 1400), stage([p95], "wipe-right", 1400, delay=-1400),
             stage([legend], delay=-300)],
            [stage([stat, stat_t])]])

# 9 ------------------------------------------------------------------ conclusion: the path
s = deck.slide("content", "Conclusion: from two sources to one platform", morph=False)
fade(s)
kicker(s, "Conclusion")
start = [box(s, A.x, 1.95, 1.7, 2.4, fill=P.bad.tint, radius=0.08),
         text(s, A.x, 3.1, 1.7, 1.0, [[("Spreadsheet", {"bold": True, "size": 12.5})],
                                       [("+ paper", {"bold": True, "size": 12.5})]], align="c"),
         icon(s, "table_view", A.x + 0.55, 2.2, 0.6, color=P.bad.dark)]
path_ = [("Map", REQ), ("Model", DATA), ("Build", APP), ("Test", None)]
chev = []
for k, (t, tone) in enumerate(path_):
    x = 2.65 + k * 2.05
    sh = s.shapes.add_shape(MSO_SHAPE.CHEVRON if k else MSO_SHAPE.PENTAGON, Inches(x), Inches(2.75), Inches(2.0),
                            Inches(0.8))
    sh.adjustments[0] = 0.28
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(tone.base if tone else P.slate)
    sh.line.fill.background()
    label(sh, t, size=15, bold=True, color=P.white)
    chev.append(sh)
end = [box(s, 10.9, 1.95, 1.7, 2.4, fill=P.good.tint, radius=0.08),
       icon(s, "devices", 11.4, 2.2, 0.7, color=P.good.dark),
       text(s, 10.9, 3.1, 1.7, 1.0, [[("One platform", {"bold": True, "size": 12.5})]], align="c")]
before = pill(s, A.x, 4.9, 3.6, 0.6, "Before: one book, two records", fill=P.bad.tint, color=P.bad.dark, size=14,
              bold=False)
arrow_ = line(s, A.x + 3.75, 5.2, A.x + 4.45, 5.2, color=P.slate, width=2.25, arrow=True)
now = pill(s, A.x + 4.6, 4.9, 4.2, 0.6, "Now: one book, one record", fill=P.good.base, color=P.white, size=16)
animate(s, [[stage(start)]] + [[stage([c], "wipe-right", 400)] for c in chev] +
        [[stage(end)], [stage([before]), stage([arrow_], "wipe-right", 350), stage([now], dur=400)]])

deck.save(HERE / "demo.pptx")
