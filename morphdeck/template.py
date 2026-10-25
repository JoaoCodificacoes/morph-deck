"""Any .pptx template, described by a small TOML file.

    [template]
    file = "neutral.pptx"          # relative to this config

    [layouts]                      # name -> layout index (see `morphdeck inspect`)
    cover = 0
    content = 1
    end = 2

    [placeholders]
    title = 0                      # idx of the title placeholder
    slide_number = 12              # idx of the slide-number placeholder, if the template has one
    keep = []                      # other placeholder idx to keep; every other placeholder is removed

    [area]                         # the content area of a content slide, in inches
    x = 0.7
    y = 1.7
    w = 11.9
    h = 4.9

    [fonts]
    body = "Aptos"
    head = "Aptos Display"
    title_size = 30

    [palette]                      # optional overrides of morphdeck.palette.Palette fields
    ink = "1F2933"

    [identities]                   # one colour per thread of the talk: base, tint, dark
    data = ["009DE0", "E3F4FC", "0077B6"]
"""
import tomllib
from dataclasses import dataclass, fields
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

from .draw import style_runs
from .morph import morph as _morph
from .palette import Palette, Tone
from .theme import Theme, use


@dataclass
class Area:
    x: float
    y: float
    w: float
    h: float

    @property
    def right(self):
        return self.x + self.w

    @property
    def bottom(self):
        return self.y + self.h


class Deck:
    """A deck built on a template. ``deck.slide("content", "Title")`` adds a slide whose title placeholder is
    named ``!!title``, so titles glide from slide to slide under Morph."""

    def __init__(self, config):
        config = Path(config)
        self.cfg = tomllib.loads(config.read_text())
        self.root = config.parent
        self.prs = Presentation(self.root / self.cfg["template"]["file"])
        ids = self.prs.slides._sldIdLst
        for s in list(ids):  # start from the template's layouts only
            self.prs.part.drop_rel(s.rId)
            ids.remove(s)
        a = self.cfg.get("area", {})
        self.area = Area(a.get("x", 0.7), a.get("y", 1.7), a.get("w", 11.9), a.get("h", 4.9))
        f = self.cfg.get("fonts", {})
        pal = Palette(**{k: v for k, v in self.cfg.get("palette", {}).items()
                         if k in {fl.name for fl in fields(Palette)} and k != "identities"})
        for name, (b, t, d) in self.cfg.get("identities", {}).items():
            pal.identities[name] = Tone(b, t, d)
        self.theme = use(Theme(pal, f.get("body", "Aptos"), f.get("head", "Aptos Display")))
        self.title_size = f.get("title_size", 30)
        self.width = self.prs.slide_width / 914400
        self.height = self.prs.slide_height / 914400

    @property
    def palette(self):
        return self.theme.palette

    def layout(self, name):
        return self.prs.slide_layouts[self.cfg["layouts"][name]]

    def slide(self, layout="content", title=None, morph=True, title_size=None):
        """New slide. Unused placeholders are removed, the title is styled with explicit fonts, and the slide is
        entered with Morph unless ``morph=False``."""
        ph_cfg = self.cfg.get("placeholders", {})
        t_idx, n_idx = ph_cfg.get("title", 0), ph_cfg.get("slide_number")
        keep = set(ph_cfg.get("keep", []))
        s = self.prs.slides.add_slide(self.layout(layout))
        for ph in list(s.placeholders):
            idx = ph.placeholder_format.idx
            if idx == t_idx and title is not None:
                ph.text = title
                ph.name = "!!title"
                style_runs(ph.text_frame, size=title_size or self.title_size, color=self.palette.ink,
                           font=self.theme.font_head,
                           bold=False)
            elif idx not in keep:
                ph._element.getparent().remove(ph._element)
        if n_idx is not None:
            for lph in s.slide_layout.placeholders:
                if lph.placeholder_format.idx == n_idx:
                    s.shapes.clone_placeholder(lph)
                    self._number(s, n_idx)
        if morph and len(self.prs.slides) > 1:
            _morph(s)
        return s

    def _number(self, s, idx):
        ph = next(p for p in s.placeholders if p.placeholder_format.idx == idx)
        p = ph.text_frame.paragraphs[0]._p
        fld = etree.SubElement(p, qn("a:fld"))
        fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}")
        fld.set("type", "slidenum")
        r_pr = etree.SubElement(fld, qn("a:rPr"))
        etree.SubElement(r_pr, qn("a:latin")).set("typeface", self.theme.font)
        etree.SubElement(fld, qn("a:t")).text = "‹#›"

    def fill(self, slide, idx, content, size=16, color=None, bold=False, font=None):
        """Put text in a kept placeholder (e.g. on the cover), always with an explicit font."""
        ph = next(p for p in slide.placeholders if p.placeholder_format.idx == idx)
        ph.text = content
        style_runs(ph.text_frame, size=size, color=color or self.palette.ink, font=font or self.theme.font,
                   bold=bold)
        return ph

    def save(self, path, lint=True):
        """Save, then run ``morphdeck lint`` and print what it finds. Returns the lint findings."""
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(path)
        if not lint:
            return []
        from .lint import lint as run_lint, report
        found = run_lint(path)
        print(report(found, path))
        return found
