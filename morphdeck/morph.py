"""Morph transitions and the rules that make them work.

Morph moves, resizes, recolours and rotates a shape from slide N to slide N+1 when both shapes have the same
name starting with ``!!``. The rules learned the hard way:

1. **The target must be static.** A shape on slide N+1 with an entrance animation is invisible when Morph runs,
   so the link breaks. Build *other* things on N+1 by click, never the shape that arrives by Morph.
2. **Morph uses resting positions.** It starts from where the shape sits on slide N, not where an animation left
   it. A shape moved by a motion path must *rest* at the end of its path (paths relative to the end point).
3. **One source, one target.** To send one shape to two places, draw a twin (``twin()``) on slide N.
4. **Names are unique per slide.** Two shapes with the same ``!!`` name make Morph guess.
5. **z-order travels.** A shape that must stay on top while it moves must be on top on both slides
   (``to_front()``).
6. **Shapes inside groups don't link reliably.** Keep the shape that should travel outside the group.

``morphdeck lint`` checks 1, 2, 4 and reports one-sided names.
"""
from copy import deepcopy

from lxml import etree
from pptx.oxml.ns import qn

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _insert_transition(slide, xml):
    sld = slide._element
    for old in sld.findall(qn("p:transition")):
        sld.remove(old)
    for old in sld.findall(f"{{{MC_NS}}}AlternateContent"):
        sld.remove(old)
    el = etree.fromstring(xml)
    clr = sld.find(qn("p:clrMapOvr"))
    (clr if clr is not None else sld.find(qn("p:cSld"))).addnext(el)


def morph(slide, dur=900):
    """Enter this slide with Morph (falls back to a fade in viewers without it)."""
    _insert_transition(slide, (
        f'<mc:AlternateContent xmlns:mc="{MC_NS}">'
        '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
        f'<p:transition xmlns:p="{P_NS}" xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" '
        f'spd="slow" p14:dur="{dur}"><p159:morph option="byObject"/></p:transition></mc:Choice>'
        f'<mc:Fallback><p:transition xmlns:p="{P_NS}" spd="slow"><p:fade/></p:transition></mc:Fallback>'
        '</mc:AlternateContent>'))


def fade(slide):
    _insert_transition(slide, f'<p:transition xmlns:p="{P_NS}" spd="med"><p:fade/></p:transition>')


def link(shape, name):
    """Name a shape so Morph matches it across slides. ``name`` without the ``!!`` prefix."""
    shape.name = "!!" + name.removeprefix("!!")
    return shape


def twin(slide, shape, name):
    """Copy ``shape`` in place under another ``!!`` name, so one shape can morph into two targets."""
    el = deepcopy(shape._element)
    tree = slide.shapes._spTree
    tree.insert(list(tree).index(shape._element), el)  # under the original, so it looks like one shape
    ids = [int(e.get("id")) for e in tree.iter() if e.tag == qn("p:cNvPr") and e.get("id", "").isdigit()]
    nv = el.find(".//" + qn("p:cNvPr"))
    new_id = max(ids) + 1
    nv.set("id", str(new_id))
    nv.set("name", "!!" + name.removeprefix("!!"))
    return next(s for s in slide.shapes if s.shape_id == new_id)


def to_front(slide, shapes):
    """Move shapes to the top of the z-order (and keep them there while Morph moves them)."""
    tree = slide.shapes._spTree
    for sh in shapes:
        tree.remove(sh._element)
        tree.append(sh._element)
