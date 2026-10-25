"""Drawing helpers on top of python-pptx. Units are inches; colours are hex strings without '#'.

Every helper takes ``name=``: give a shape a name starting with ``!!`` and Morph will carry it to the shape with
the same name on the next slide (see ``morphdeck.morph``).
"""
from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .theme import current


def rgb(h):
    return RGBColor.from_string(h)


def _pal():
    return current().palette


def style_runs(tf, size=None, color=None, font=None, bold=None):
    """Restyle every run of a text frame. Pass ``font`` for placeholder text: a template font left implicit may
    be missing on the presenting machine, and PowerPoint then spaces the letters badly."""
    for p in tf.paragraphs:
        for r in p.runs:
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = rgb(color)
            if font:
                r.font.name = font
            if bold is not None:
                r.font.bold = bold


def box(slide, x, y, w, h, fill=None, line=None, radius=0.08, name=None, shape=None, lw=1.0):
    """Rectangle (rounded when ``radius`` > 0). ``fill=""`` for no fill; default fill is the card colour."""
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    fill = _pal().card if fill is None else fill
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    if name:
        sh.name = name
    sh.text_frame.text = ""
    return sh


def text(slide, x, y, w, h, content, size=14, color=None, bold=False, align="l", anchor="t", font=None,
         name=None, spacing=None, shape=None):
    """Text box. ``content``: a string, or a list of paragraphs; a paragraph is a string or a list of runs, and a
    run is a string or ``(text, overrides)`` with overrides among size, color, bold, italic, font."""
    th = current()
    color = color or th.palette.ink
    font = font or th.font
    if shape is None:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shape.name = name
    tf = shape.text_frame
    tf.word_wrap = True
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(0.04))
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if spacing:
            p.space_after = Pt(spacing)
        runs = para if isinstance(para, list) else [(para, {})]
        for run in runs:
            if isinstance(run, str):
                run = (run, {})
            t, o = run
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)  # always explicit, never inherited from the template
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = rgb(o.get("color", color))
    return shape


def label(sh, content, size=14, color=None, bold=False, align="c", anchor="m", **kw):
    """Write text inside an existing shape."""
    return text(None, 0, 0, 0, 0, content, size=size, color=color, bold=bold, align=align, anchor=anchor,
                shape=sh, **kw)


def _ends(ln, arrow, dash):
    # schema order inside a:ln: prstDash must come before headEnd/tailEnd, or PowerPoint ignores it
    if dash:
        etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    if arrow:
        t = etree.SubElement(ln, qn("a:tailEnd"))
        t.set("type", "triangle")
        t.set("w", "med")
        t.set("h", "med")


def line(slide, x1, y1, x2, y2, color=None, width=1.5, name=None, arrow=False, dash=False):
    """Straight connector from (x1, y1) to (x2, y2); ``arrow`` puts a head at the end."""
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color or _pal().line)
    c.line.width = Pt(width)
    _ends(c.line._get_or_add_ln(), arrow, dash)
    if name:
        c.name = name
    return c


def polyline(slide, pts, color=None, width=2.5, name=None, arrow=False, dash=False):
    """Open polyline through (x, y) points."""
    fb = slide.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
    fb.add_line_segments([(Inches(x), Inches(y)) for x, y in pts[1:]], close=False)
    sh = fb.convert_to_shape()
    sh.fill.background()
    sh.line.color.rgb = rgb(color or _pal().ink)
    sh.line.width = Pt(width)
    sh.shadow.inherit = False
    _ends(sh.line._get_or_add_ln(), arrow, dash)
    if name:
        sh.name = name
    return sh


def picture(slide, path, x, y, w=None, h=None, name=None, crop=None):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None,
                                   Inches(h) if h else None)
    if crop:
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
    if name:
        pic.name = name
    return pic


def round_corners(pic, adj=0.08):
    """Give a picture rounded-rectangle geometry."""
    geom = pic._element.spPr.find(qn("a:prstGeom"))
    geom.set("prst", "roundRect")
    av = geom.find(qn("a:avLst"))
    if av is None:
        av = etree.SubElement(geom, qn("a:avLst"))
    for gd in list(av):
        av.remove(gd)
    gd = etree.SubElement(av, qn("a:gd"))
    gd.set("name", "adj")
    gd.set("fmla", f"val {int(adj * 100000)}")


def phone(slide, path, x, y, h, aspect=780 / 1688, name=None):
    """A screenshot inside a phone bezel, ``h`` inches tall. Returns ``[bezel, screen]``."""
    bez = 0.07
    sw = (h - 2 * bez) * aspect
    frame = box(slide, x, y, sw + 2 * bez, h, fill=_pal().navy, radius=0.1, name=f"{name}-frame" if name else None)
    pic = picture(slide, path, x + bez, y + bez, w=sw, h=h - 2 * bez, name=f"{name}-screen" if name else None)
    round_corners(pic, 0.085)
    return [frame, pic]


def pill(slide, x, y, w, h, content, fill=None, color=None, size=12, bold=True, name=None, line_c=None):
    """Fully rounded label, e.g. a tag or a status."""
    p = _pal()
    sh = box(slide, x, y, w, h, fill=fill or p.card, radius=0.5, name=name, line=line_c)
    label(sh, content, size=size, color=color or p.ink, bold=bold)
    return sh


def cylinder(slide, x, y, w, h, fill=None, name=None):
    sh = box(slide, x, y, w, h, fill=fill, shape=MSO_SHAPE.CAN, radius=0, name=name)
    sh.adjustments[0] = 0.18
    return sh


def shadow(sh, blur=8, dist=3, alpha=22):
    """Soft drop shadow. Reuses the shape's effect list: a second one makes PowerPoint refuse the file."""
    sp = sh._element.spPr
    eff = sp.find(qn("a:effectLst"))
    if eff is None:
        eff = etree.SubElement(sp, qn("a:effectLst"))
    o = etree.SubElement(eff, qn("a:outerShdw"))
    o.set("blurRad", str(int(Pt(blur))))
    o.set("dist", str(int(Pt(dist))))
    o.set("dir", "5400000")
    o.set("algn", "t")
    o.set("rotWithShape", "0")
    c = etree.SubElement(o, qn("a:srgbClr"))
    c.set("val", "000000")
    etree.SubElement(c, qn("a:alpha")).set("val", str(alpha * 1000))


def dashed(sh):
    """Dashed outline on an existing shape."""
    ln = sh.line._get_or_add_ln()
    etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    return sh
