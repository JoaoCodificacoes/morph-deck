"""morphdeck — command line.

    morphdeck inspect template.pptx [-o template.toml]   layouts, placeholders and fonts; writes a starter config
    morphdeck lint deck.pptx [-v]                       checks for broken Morph links and animations
    morphdeck render deck.pptx out/                     slide PNGs (PowerPoint on Windows or WSL)
    morphdeck video deck.pptx --slides 3,4 -o v.mp4     record slides, clicks auto-advance (PowerPoint)
    morphdeck sheet v.mp4 [-o sheet.png]                contact sheet of frames, to look at motion as one image
    morphdeck gif v.mp4 -o clip.gif [--start 2 --end 5] short GIF, e.g. for a README
"""
import argparse
import platform
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

PS_DIR = Path(__file__).parent / "ps"


# ---------------------------------------------------------------- PowerPoint bridge (Windows or WSL)

def _wsl():
    return platform.system() == "Linux" and shutil.which("powershell.exe") is not None


def _win_temp():
    """A working folder Windows PowerPoint can open files from."""
    if platform.system() == "Windows":
        d = Path(tempfile.gettempdir()) / "morphdeck"
    elif _wsl():
        t = subprocess.run(["cmd.exe", "/c", "echo %TEMP%"], capture_output=True, text=True, cwd="/mnt/c").stdout
        d = Path(subprocess.run(["wslpath", "-u", t.strip()], capture_output=True, text=True).stdout.strip()) / "morphdeck"
    else:
        sys.exit("rendering needs PowerPoint: run on Windows, or in WSL with PowerPoint installed on Windows")
    d.mkdir(parents=True, exist_ok=True)
    return d


def _winpath(p):
    if platform.system() == "Windows":
        return str(p)
    return subprocess.run(["wslpath", "-w", str(p)], capture_output=True, text=True).stdout.strip()


def _ps(script, *args):
    exe = "powershell" if platform.system() == "Windows" else "powershell.exe"
    tmp = _win_temp()
    local = tmp / script  # PowerShell runs scripts from Windows paths more reliably than from \\wsl$
    shutil.copy(PS_DIR / script, local)
    return subprocess.run([exe, "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", _winpath(local), *args])


def render(deck, out, width=1280):
    tmp = _win_temp()
    src, dst = tmp / "render.pptx", tmp / "render_out"
    shutil.copy(deck, src)
    shutil.rmtree(dst, ignore_errors=True)
    r = _ps("render.ps1", _winpath(src), _winpath(dst), str(width))
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    for old in out.glob("s*.png"):
        old.unlink()
    for png in sorted(dst.glob("*.png")):
        shutil.copy(png, out / png.name)
    return r.returncode


def video(deck, out, slides=None, advance=1):
    tmp = _win_temp()
    src, dst = tmp / "video.pptx", tmp / "video.mp4"
    shutil.copy(deck, src)
    dst.unlink(missing_ok=True)
    r = _ps("video.ps1", _winpath(src), _winpath(dst), slides or "", str(advance))
    if r.returncode or not dst.exists():
        print("video export failed; close the deck in PowerPoint if it is open for editing, then retry")
        return 1
    shutil.copy(dst, out)
    return 0


# ---------------------------------------------------------------- frames

def _ffmpeg():
    exe = shutil.which("ffmpeg")
    if not exe:
        sys.exit("ffmpeg not found")
    return exe


def sheet(vid, out, fps=6, cols=8, rows=6, width=256, start=0.0):
    vf = f"fps={fps},scale={width}:-2,tile={cols}x{rows}"
    return subprocess.run([_ffmpeg(), "-v", "error", "-y", "-ss", str(start), "-i", str(vid), "-vf", vf,
                           "-frames:v", "1", str(out)]).returncode


def gif(vid, out, start=0.0, end=None, width=720, fps=15):
    span = ["-ss", str(start)] + (["-to", str(end)] if end else [])
    vf = f"fps={fps},scale={width}:-1:flags=lanczos,split[a][b];[a]palettegen=stats_mode=diff[p];[b][p]paletteuse"
    return subprocess.run([_ffmpeg(), "-v", "error", "-y", *span, "-i", str(vid), "-vf", vf, str(out)]).returncode


# ---------------------------------------------------------------- inspect

def inspect(template, out=None):
    import re
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER as PH

    prs = Presentation(str(template))
    emu = 914400
    print(f"{template}: {prs.slide_width / emu:.3f} x {prs.slide_height / emu:.3f} in")
    theme = next((p for p in prs.part.package.iter_parts() if "theme" in str(p.partname)), None)
    fonts = {}
    if theme is not None:
        xml = theme.blob.decode("utf8", "ignore")
        for kind in ("major", "minor"):
            m = re.search(rf"<a:{kind}Font><a:latin typeface=\"([^\"]*)\"", xml)
            fonts[kind] = m.group(1) if m else ""
        print(f"theme fonts: headings '{fonts['major']}', body '{fonts['minor']}'")
    guess = {}
    for i, lay in enumerate(prs.slide_layouts):
        phs = [(p.placeholder_format.idx, p.placeholder_format.type) for p in lay.placeholders]
        desc = ", ".join(f"{idx}:{str(t).split('.')[-1].split(' ')[0].lower()}" for idx, t in phs)
        print(f"  layout {i:2}  {lay.name:32} {desc}")
        types = {t for _, t in phs}
        if "cover" not in guess and PH.CENTER_TITLE in types:
            guess["cover"] = i
        if PH.TITLE in types:  # content: a title and a slide number beats a title only, beats a title + body
            extra = len([t for _, t in phs if t not in (PH.TITLE, PH.FOOTER, PH.DATE, PH.SLIDE_NUMBER)])
            rank = (0 if PH.SLIDE_NUMBER in types else 1, extra)
            if rank < guess.get("_rank", (9, 99)):
                guess["content"], guess["_rank"] = i, rank
        for idx, t in phs:
            if t == PH.SLIDE_NUMBER:
                guess.setdefault("slide_number", idx)
    guess.setdefault("cover", 0)
    guess.setdefault("content", min(1, len(prs.slide_layouts) - 1))
    guess.setdefault("end", len(prs.slide_layouts) - 1)
    lay = prs.slide_layouts[guess["content"]]
    title = next((p for p in lay.placeholders if p.placeholder_format.type == PH.TITLE), None)
    w, h = prs.slide_width / emu, prs.slide_height / emu
    top = (title.top + title.height) / emu + 0.25 if title is not None else 1.7
    left = title.left / emu if title is not None else 0.7
    if not (0.5 < top < h * 0.45):  # title placeholder in an odd place: fall back to a common layout
        top = 1.7
    if not (0.2 < left < w * 0.2):
        left = 0.7
    config = f'''[template]
file = "{Path(template).name}"

[layouts]
cover = {guess["cover"]}
content = {guess["content"]}
end = {guess["end"]}

[placeholders]
title = 0
{"slide_number = " + str(guess["slide_number"]) if "slide_number" in guess else "# slide_number = 12"}
keep = []

[area]
x = {left:.2f}
y = {top:.2f}
w = {w - 2 * left:.2f}
h = {h - top - 0.9:.2f}

[fonts]
body = "{fonts.get("minor") or "Aptos"}"
head = "{fonts.get("major") or "Aptos Display"}"
title_size = 30

[identities]
# one colour per thread of the talk: base, tint, dark
# example = ["009DE0", "E3F4FC", "0077B6"]
'''
    if out:
        Path(out).write_text(config)
        print(f"wrote {out} — check the layout guesses above before using it")
    else:
        print("\n" + config)
    return 0


def main(argv=None):
    ap = argparse.ArgumentParser(prog="morphdeck", description=__doc__.split("\n\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("inspect")
    p.add_argument("template")
    p.add_argument("-o", "--out")
    p = sub.add_parser("lint")
    p.add_argument("deck")
    p.add_argument("-v", "--verbose", action="store_true")
    p = sub.add_parser("render")
    p.add_argument("deck")
    p.add_argument("out")
    p.add_argument("--width", type=int, default=1280)
    p = sub.add_parser("video")
    p.add_argument("deck")
    p.add_argument("--slides", help="comma-separated slide numbers, e.g. 3,4")
    p.add_argument("-o", "--out", default="video.mp4")
    p.add_argument("--advance", type=int, default=1, help="seconds before each click")
    p = sub.add_parser("sheet")
    p.add_argument("video")
    p.add_argument("-o", "--out", default="sheet.png")
    p.add_argument("--fps", type=float, default=6)
    p.add_argument("--start", type=float, default=0)
    p = sub.add_parser("gif")
    p.add_argument("video")
    p.add_argument("-o", "--out", default="clip.gif")
    p.add_argument("--start", type=float, default=0)
    p.add_argument("--end", type=float)
    a = ap.parse_args(argv)
    if a.cmd == "inspect":
        return inspect(a.template, a.out)
    if a.cmd == "lint":
        from .lint import main as lint_main
        return lint_main([a.deck] + (["-v"] if a.verbose else []))
    if a.cmd == "render":
        return render(a.deck, a.out, a.width)
    if a.cmd == "video":
        return video(a.deck, a.out, a.slides, a.advance)
    if a.cmd == "sheet":
        return sheet(a.video, a.out, a.fps, start=a.start)
    if a.cmd == "gif":
        return gif(a.video, a.out, a.start, a.end)


if __name__ == "__main__":
    sys.exit(main())
